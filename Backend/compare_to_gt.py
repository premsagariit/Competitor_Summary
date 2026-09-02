"""Compare Data_Engine_UI.xlsx against the ground-truth GT_Data_Engine.xlsx."""
import openpyxl
from collections import Counter, defaultdict

HEADERS = ["Slide #", "Category", "Company", "Meric 1", "Metric 2",
           "Source Tab", "Link to Source document", "FY26_Q3", "FY25_Q3", "Growth"]
COL = {h: i + 1 for i, h in enumerate(HEADERS)}


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    return wb["Data Engine"]


def row_key(ws, r):
    return (
        ws.cell(row=r, column=COL["Slide #"]).value,
        ws.cell(row=r, column=COL["Company"]).value,
        ws.cell(row=r, column=COL["Meric 1"]).value,
        ws.cell(row=r, column=COL["Metric 2"]).value,
    )


def num_close(a, b, rel=0.01, abs_tol=0.05):
    if a is None or b is None:
        return a is None and b is None
    if not isinstance(a, (int, float)) or not isinstance(b, (int, float)):
        return a == b
    if a == b:
        return True
    denom = max(abs(a), abs(b), 1e-9)
    return abs(a - b) / denom <= rel or abs(a - b) <= abs_tol


def main():
    mine = load("Data_Engine_UI.xlsx")
    gt = load("GT_Data_Engine.xlsx")

    gt_rows = {}
    for r in range(2, gt.max_row + 1):
        slide = gt.cell(row=r, column=COL["Slide #"]).value
        if not isinstance(slide, int):
            continue
        gt_rows[row_key(gt, r)] = r

    mine_rows = {}
    for r in range(2, mine.max_row + 1):
        slide = mine.cell(row=r, column=COL["Slide #"]).value
        if not isinstance(slide, int):
            continue
        mine_rows[row_key(mine, r)] = r

    unmatched_keys = set(mine_rows) - set(gt_rows)
    print(f"Row-key alignment: {len(mine_rows)} rows in mine, {len(gt_rows)} in GT, "
          f"{len(set(mine_rows) & set(gt_rows))} keys match, {len(unmatched_keys)} of mine have no GT key match")
    if unmatched_keys:
        print("Sample unmatched keys (first 10):")
        for k in list(unmatched_keys)[:10]:
            print("  ", k)

    total_gt_filled = 0
    both_filled_correct = 0
    both_filled_wrong = 0
    mine_filled_gt_blank = 0
    gt_filled_mine_blank = 0
    both_blank = 0

    by_company = defaultdict(lambda: Counter())
    wrong_examples = []
    blank_examples = []

    for key, gr in gt_rows.items():
        gt_cur = gt.cell(row=gr, column=COL["FY26_Q3"]).value
        gt_prior = gt.cell(row=gr, column=COL["FY25_Q3"]).value
        gt_has = gt_cur is not None or gt_prior is not None

        mr = mine_rows.get(key)
        if mr is None:
            mine_cur = mine_prior = None
        else:
            mine_cur = mine.cell(row=mr, column=COL["FY26_Q3"]).value
            mine_prior = mine.cell(row=mr, column=COL["FY25_Q3"]).value
        mine_has = mine_cur is not None or mine_prior is not None

        company = key[1]
        slide = key[0]

        if gt_has:
            total_gt_filled += 1
        if not gt_has and not mine_has:
            both_blank += 1
            continue
        if gt_has and not mine_has:
            gt_filled_mine_blank += 1
            by_company[company]["missing"] += 1
            if len(blank_examples) < 30:
                blank_examples.append((slide, *key[1:], gt_cur, gt_prior))
            continue
        if mine_has and not gt_has:
            mine_filled_gt_blank += 1
            by_company[company]["extra"] += 1
            continue

        cur_ok = num_close(mine_cur, gt_cur)
        prior_ok = num_close(mine_prior, gt_prior)
        if cur_ok and prior_ok:
            both_filled_correct += 1
            by_company[company]["correct"] += 1
        else:
            both_filled_wrong += 1
            by_company[company]["wrong"] += 1
            if len(wrong_examples) < 40:
                wrong_examples.append((slide, *key[1:], (mine_cur, mine_prior), (gt_cur, gt_prior)))

    print()
    print("===== SUMMARY =====")
    print(f"GT rows with a value filled (denominator): {total_gt_filled}")
    print(f"  Correctly extracted & matching GT:        {both_filled_correct}")
    print(f"  Filled but WRONG (mismatch vs GT):         {both_filled_wrong}")
    print(f"  GT has a value, we left blank (missed):    {gt_filled_mine_blank}")
    print(f"We filled but GT is blank (extra/unverified): {mine_filled_gt_blank}")
    print(f"Both blank: {both_blank}")
    accuracy = both_filled_correct / total_gt_filled * 100 if total_gt_filled else 0
    coverage = (both_filled_correct + both_filled_wrong) / total_gt_filled * 100 if total_gt_filled else 0
    print(f"\nAccuracy (correct / GT-filled): {accuracy:.1f}%")
    print(f"Coverage (attempted / GT-filled): {coverage:.1f}%")

    print()
    print("===== BY COMPANY =====")
    for company, c in sorted(by_company.items(), key=lambda x: -sum(x[1].values())):
        tot = sum(c.values())
        print(f"  {company or '(blank)'}: correct={c['correct']} wrong={c['wrong']} missing={c['missing']} extra={c['extra']}  (n={tot})")

    print()
    print("===== WRONG EXAMPLES (first 40) =====")
    for slide, company, metric1, metric2, mine_v, gt_v in wrong_examples:
        print(f"  slide{slide} {company} / {metric1} / {metric2}: mine={mine_v} gt={gt_v}")

    print()
    print("===== MISSING EXAMPLES (GT filled, we're blank) (first 30) =====")
    for slide, company, metric1, metric2, gt_cur, gt_prior in blank_examples:
        print(f"  slide{slide} {company} / {metric1} / {metric2}: gt=({gt_cur},{gt_prior})")


if __name__ == "__main__":
    main()
