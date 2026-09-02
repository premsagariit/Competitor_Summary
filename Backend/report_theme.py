"""
Shared visual theme for the FY26 Q3 Competition Analysis deck: slide
dimensions, company color palette, company-name resolution (the Data Engine
uses several inconsistent spellings for the same company across different
slides - GIC full legal names on Slides 3-9, short pipeline keys on
Slides 13-35, and a third short form as Slide 11/12's own metric1 values),
and the reusable chrome (header/footer/title-badge/insight-panel) drawn on
every slide, matching Backend/Competition Summary FY25 1.pdf's layout.
"""
import os

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

SLIDE_WIDTH = Inches(8.27)   # A4 portrait, matching the reference PDF's page size
SLIDE_HEIGHT = Inches(11.69)

LOGO_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "logos")

# ---------------------------------------------------------------------------
# Colors
# ---------------------------------------------------------------------------

NAVY = RGBColor(0x1F, 0x35, 0x64)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
CYAN = RGBColor(0x00, 0xB0, 0xF0)
ORANGE = RGBColor(0xED, 0x7D, 0x31)
CREAM = RGBColor(0xFF, 0xF7, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x33, 0x33, 0x33)
LIGHT_GREY = RGBColor(0xF2, 0xF2, 0xF2)

# Per-company colors - the 5 the reference deck itself uses (sampled from
# its charts), plus 2 new ones for companies it never plots (Narayana,
# Galaxy - our FY26 Q3 pipeline extracts both, the reference deck doesn't).
COMPANY_COLORS = {
    "NBHI": RGBColor(0x9D, 0xC3, 0xE6),
    "STAR": RGBColor(0x1F, 0x4E, 0x79),
    "CARE": RGBColor(0xFF, 0xD9, 0x1A),
    "CIGNA": RGBColor(0x92, 0xD0, 0x50),
    "ABHI": RGBColor(0xC0, 0x00, 0x00),
    "NARAYANA": RGBColor(0x2C, 0xA6, 0xA4),
    "GALAXY": RGBColor(0xED, 0x7D, 0x31),
}
COMPANY_DISPLAY_NAME = {
    "NBHI": "NBHI", "STAR": "STAR", "CARE": "CARE", "CIGNA": "CIGNA",
    "ABHI": "ABHI", "NARAYANA": "Narayana", "GALAXY": "Galaxy",
}
# Company keys that have a real extracted logo image (Backend/assets/logos).
COMPANIES_WITH_LOGO = {"NBHI", "STAR", "CARE", "CIGNA", "ABHI"}

# Market-segment colors (Private/Public/SAHI/Specialized etc.), reused across
# Slides 3-6's industry-level doughnut charts - independent of company colors.
SEGMENT_COLORS = {
    "Private": BLUE, "Pvt GI": BLUE, "Pvt. GI": BLUE,
    "Public": RGBColor(0xFF, 0xC0, 0x00), "Public GI": RGBColor(0xFF, 0xC0, 0x00),
    "SAHI": RGBColor(0x70, 0xAD, 0x47), "SAHI Market": RGBColor(0x70, 0xAD, 0x47),
    "Specialized Insurer": NAVY, "Specialised": NAVY,
    "Industry": BLUE,
}
FALLBACK_SERIES_COLORS = [
    BLUE, ORANGE, RGBColor(0x70, 0xAD, 0x47), RGBColor(0xFF, 0xC0, 0x00),
    RGBColor(0x7C, 0x3A, 0xED), RGBColor(0x26, 0x4E, 0x70), RGBColor(0xC0, 0x00, 0x00),
    RGBColor(0x00, 0xB0, 0xB0), RGBColor(0xA9, 0xD1, 0x8E), RGBColor(0xBF, 0xBF, 0xBF),
]


def series_colors_for(names):
    """{name: color}, cycling FALLBACK_SERIES_COLORS - used for stacked-chart
    series that don't have a fixed semantic color (asset classes, states,
    channels, intermediary types)."""
    return {name: FALLBACK_SERIES_COLORS[i % len(FALLBACK_SERIES_COLORS)] for i, name in enumerate(names)}


def canonical_company(raw):
    """Resolves any of the Data Engine's several spellings for a company
    down to one of the 7 short keys in COMPANY_COLORS, or None if `raw`
    isn't an individual company at all (e.g. 'Industry', 'SAHI', 'Public GI'
    aggregate/group labels used as the Company or Meric-1 value on several
    slides)."""
    if not raw:
        return None
    t = str(raw).strip().lower()
    if "niva" in t or t == "nbhi":
        return "NBHI"
    if "star" in t:
        return "STAR"
    if "care" in t:
        return "CARE"
    if "cigna" in t or "manipal" in t:
        return "CIGNA"
    if "aditya" in t or t == "abhi":
        return "ABHI"
    if "narayana" in t or "naryana" in t:  # "Naryana" is a known sheet typo (Slide 5)
        return "NARAYANA"
    if "galaxy" in t:
        return "GALAXY"
    return None


def logo_path(company_key):
    if company_key not in COMPANIES_WITH_LOGO:
        return None
    p = os.path.join(LOGO_DIR, f"{company_key}.png")
    return p if os.path.exists(p) else None


# ---------------------------------------------------------------------------
# Slide chrome
# ---------------------------------------------------------------------------

def new_slide(prs, layout_idx=6):
    """layout 6 is the standard 'blank' layout in the default template."""
    return prs.slides.add_slide(prs.slide_layouts[layout_idx])


def _no_line(shape):
    shape.line.fill.background()


def add_header(slide, subtitle="Competition Analysis FY26 Q3"):
    tb = slide.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(4), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = subtitle
    run.font.size = Pt(11)
    run.font.color.rgb = NAVY
    run.font.bold = False

    logo = logo_path("NBHI")
    if logo:
        slide.shapes.add_picture(logo, Inches(6.9), Inches(0.15), height=Inches(0.45))

    line = slide.shapes.add_connector(1, Inches(0.3), Inches(0.62), Inches(7.97), Inches(0.62))
    line.line.color.rgb = CYAN
    line.line.width = Pt(2.25)


def add_footer(slide, page_number):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(11.25), SLIDE_WIDTH, Inches(0.44))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CYAN
    _no_line(bar)
    tf = bar.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = 'Our Vision is, "To become India’s most admired Health Insurance Company"'
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = WHITE

    tab = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0), Inches(11.0), Inches(1.0), Inches(0.32))
    tab.fill.solid()
    tab.fill.fore_color.rgb = RGBColor(0x8D, 0xB4, 0xE2)
    _no_line(tab)
    ptf = tab.text_frame
    pp = ptf.paragraphs[0]
    pp.alignment = PP_ALIGN.CENTER
    r = pp.add_run()
    r.text = f"Page | {page_number}"
    r.font.size = Pt(9)
    r.font.color.rgb = WHITE


def add_title_badge(slide, text):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), Inches(0.85), Inches(3.2), Inches(0.5))
    badge.fill.solid()
    badge.fill.fore_color.rgb = NAVY
    _no_line(badge)
    tf = badge.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = WHITE


def add_panel_frame(slide, left, top, width, height):
    """The dashed-orange rounded-rect frame used around every chart on the
    reference deck. Returns the shape in case a caller wants to layer inside it."""
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    frame.fill.background()
    frame.line.color.rgb = ORANGE
    frame.line.width = Pt(0.75)
    frame.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    frame.shadow.inherit = False
    return frame


def add_insight_panel(slide, bullets, left, top, width, height):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = CREAM
    panel.line.color.rgb = ORANGE
    panel.line.width = Pt(0.75)
    panel.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    panel.shadow.inherit = False
    tf = panel.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = f"➢  {bullet}"
        run.font.size = Pt(11)
        run.font.color.rgb = DARK_TEXT


def add_note(slide, text, left, top, width, height):
    """A plain 'data not available' / footnote textbox."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)


def add_placeholder_panel(slide, message, left, top, width, height):
    """Used when a slide's entire metric group has no FY26 Q3 data - an
    explicit placeholder instead of a broken/empty native chart."""
    frame = add_panel_frame(slide, left, top, width, height)
    tf = frame.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = message
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
