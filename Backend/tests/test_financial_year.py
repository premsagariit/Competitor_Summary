"""Tests for the financial-year format.

The canonical form is the SPAN — "FY25-26" means April 2025 to March 2026.
A single-year form is ambiguous (does FY26 start or end in 2026?), which
mattered most when telling the retrieval agent which year to look for.
"""
import pytest

from competitor_analysis import config as cfg


# ---------------------------------------------------------------------------
# Parsing and normalisation
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("given,expected", [
    ("FY25-26", "FY25-26"),      # canonical
    ("fy25-26", "FY25-26"),      # case
    ("FY 25-26", "FY25-26"),     # stray space
    ("25-26", "FY25-26"),        # no prefix
    ("2025-26", "FY25-26"),      # 4-digit start
    ("2025-2026", "FY25-26"),    # both 4-digit
    ("FY2025-2026", "FY25-26"),
    ("FY25_26", "FY25-26"),      # underscore
    ("FY25/26", "FY25-26"),      # slash
])
def test_accepted_spellings_normalise_to_the_span(given, expected):
    assert cfg.normalize_fy(given) == expected


@pytest.mark.parametrize("legacy,expected", [
    ("FY26", "FY25-26"),
    ("FY27", "FY26-27"),
    ("FY25", "FY24-25"),
    ("2026", "FY25-26"),
])
def test_legacy_single_year_still_works(legacy, expected):
    """The old form meant the year the FY ENDS. Existing callers, saved UI
    state and on-disk names must keep resolving to the same period."""
    assert cfg.normalize_fy(legacy) == expected


@pytest.mark.parametrize("bad", ["FY25-27", "FY26-26", "FY9", "", "banana",
                                 "FY25-24", None])
def test_invalid_years_are_rejected_with_the_expected_format(bad):
    with pytest.raises(ValueError) as e:
        cfg.normalize_fy(bad)
    msg = str(e.value)
    # The message has to tell the user what to type instead.
    assert "FY25-26" in msg or "one year boundary" in msg


def test_a_span_must_cross_exactly_one_year_boundary():
    """FY25-27 is not a financial year; catching it stops a typo becoming a
    silently wrong period."""
    with pytest.raises(ValueError, match="one year boundary"):
        cfg.normalize_fy("FY25-27")


def test_set_period_normalises():
    cfg.set_period("FY26", "Q3")          # legacy in
    assert cfg.FY == "FY25-26"            # canonical out
    cfg.set_period("2026-27", "q1")
    assert cfg.FY == "FY26-27" and cfg.QUARTER == "Q1"


# ---------------------------------------------------------------------------
# Derivations
# ---------------------------------------------------------------------------

def test_start_and_end_years():
    assert cfg.fy_start_year("FY25-26") == 2025
    assert cfg.fy_end_year("FY25-26") == 2026


@pytest.mark.parametrize("fy,quarter,month,year", [
    ("FY25-26", "Q1", "June", 2025),
    ("FY25-26", "Q2", "September", 2025),
    ("FY25-26", "Q3", "December", 2025),
    ("FY25-26", "Q4", "March", 2026),      # Q4 falls in the FY's END year
    ("FY26-27", "Q1", "June", 2026),
])
def test_calendar_mapping(fy, quarter, month, year):
    assert cfg.calendar_mapping(fy, quarter) == {"month": month, "year": year}


def test_prior_and_next_shift_the_whole_span():
    assert cfg.prior_fy("FY25-26") == "FY24-25"
    assert cfg.next_fy("FY25-26") == "FY26-27"
    # Prior period is the SAME quarter one financial year back.
    assert cfg.period_column(cfg.prior_fy("FY25-26"), "Q3") == "FY24-25_Q3"


def test_column_headers_use_the_span():
    assert cfg.period_column("FY25-26", "Q3") == "FY25-26_Q3"
    # A legacy input must not leak a legacy header into the workbook.
    assert cfg.period_column("FY26", "Q3") == "FY25-26_Q3"


def test_fy_options_are_derived_from_the_clock():
    from datetime import datetime
    # India's FY starts in April, so in March the current FY is still the one
    # that began the previous calendar year.
    march = cfg.fy_options(back=1, forward=0, today=datetime(2026, 3, 15))
    assert march[0] == "FY25-26"
    april = cfg.fy_options(back=1, forward=0, today=datetime(2026, 4, 1))
    assert april[0] == "FY26-27"
    # Newest first, and every entry is canonical.
    opts = cfg.fy_options()
    assert opts == sorted(opts, reverse=True)
    for o in opts:
        assert cfg.normalize_fy(o) == o


# ---------------------------------------------------------------------------
# The agent-facing label (why the span format exists)
# ---------------------------------------------------------------------------

def test_agent_label_comes_from_the_fy_not_the_quarter():
    """Regression: the retrieval prompt used to derive the financial year from
    the QUARTER's calendar year. For Q4 that names the following FY - a
    FY25-26 Q4 filing (March 2026) was advertised to the agent as
    'FY 2026-27', pointing it at the wrong portal dropdown."""
    assert cfg.fy_label("FY25-26") == "2025-26"
    assert cfg.fy_label_long("FY25-26") == "2025-2026"
    # Every quarter of one FY carries the same label, Q4 included.
    for q in ("Q1", "Q2", "Q3", "Q4"):
        cfg.set_period("FY25-26", q)
        assert cfg.fy_label() == "2025-26", q


def test_retrieval_prompt_states_the_right_financial_year():
    from competitor_analysis.ingestion import scraper
    for fy, quarter, expected in [("FY25-26", "Q3", "2025-26"),
                                  ("FY25-26", "Q4", "2025-26"),
                                  ("FY26-27", "Q1", "2026-27")]:
        cal = cfg.calendar_mapping(fy, quarter)
        objective, _ = scraper.build_objective(
            "Care Health Insurance", "https://example.test", cal, fy)
        assert f"Financial Year {expected}" in objective, (fy, quarter)
        # And it must not mention any other financial year.
        wrong = {"2024-25", "2025-26", "2026-27", "2027-28"} - {expected}
        for w in wrong:
            assert w not in objective, f"{fy} {quarter} prompt also mentions {w}"


# ---------------------------------------------------------------------------
# Period-derived paths
# ---------------------------------------------------------------------------

def test_paths_use_the_span_directory_name():
    assert cfg.download_dir("FY25-26", "Q3").endswith(r"downloads\FY25-26\Q3")
    assert cfg.audit_dir("FY25-26", "Q3").endswith(r"extraction_audit\FY25-26\Q3")
    assert cfg.output_pdf_path("FY25-26", "Q3").endswith(
        "Competition_Summary_FY25-26_Q3.pdf")


def test_legacy_and_canonical_resolve_to_the_same_paths():
    """A caller still passing FY26 must reach the same files as FY25-26 -
    otherwise the migration would silently split the data in two."""
    assert cfg.download_dir("FY26", "Q3") == cfg.download_dir("FY25-26", "Q3")
    assert cfg.gic_path("FY26", "Q3") == cfg.gic_path("FY25-26", "Q3")


# ---------------------------------------------------------------------------
# Period-derived state must follow set_period(), not import order
# ---------------------------------------------------------------------------

def _seed(root, quarter, companies):
    d = root / "FY25-26" / quarter
    d.mkdir(parents=True, exist_ok=True)
    for c in companies:
        (d / cfg.COMPANY_PDF_FILENAMES[c]).write_bytes(b"%PDF-1.4\n")


def test_switching_period_repoints_the_insurer_pdf_map(tmp_path, monkeypatch):
    """A second run for a different quarter, in the same process, must read
    that quarter's filings.

    pdf_cache used to resolve COMPANY_PDFS once at import, so in the
    long-lived API server a Q4 run following a Q3 run read Q3's PDFs while
    writing a Q4-labelled workbook and Q4 cache entries - wrong, and silent,
    because the numbers looked entirely plausible."""
    from competitor_analysis import paths
    from competitor_analysis.extraction import pdf_cache

    _seed(tmp_path, "Q3", ["NBHI", "ABHI"])
    _seed(tmp_path, "Q4", ["NBHI"])
    monkeypatch.setattr(paths, "DOWNLOADS_DIR", tmp_path)

    cfg.set_period("FY25-26", "Q3")
    assert set(pdf_cache.COMPANY_PDFS) == {"NBHI", "ABHI"}
    assert "Q3" in pdf_cache.COMPANY_PDFS["NBHI"]

    cfg.set_period("FY25-26", "Q4")
    assert set(pdf_cache.COMPANY_PDFS) == {"NBHI"}, "Q4 must not inherit Q3's insurers"
    assert "Q4" in pdf_cache.COMPANY_PDFS["NBHI"]


def test_the_pdf_map_is_mutated_in_place_so_importers_see_the_switch(tmp_path, monkeypatch):
    """forms, gemini and data_engine all bind the dict by name at import
    (`from ... import COMPANY_PDFS`), so refreshing it by REBINDING in
    pdf_cache would leave every one of them pointing at the old period."""
    from competitor_analysis import paths
    from competitor_analysis.extraction import pdf_cache, forms, gemini, data_engine

    _seed(tmp_path, "Q3", ["NBHI", "ABHI"])
    _seed(tmp_path, "Q4", ["NBHI"])
    monkeypatch.setattr(paths, "DOWNLOADS_DIR", tmp_path)

    cfg.set_period("FY25-26", "Q4")
    for mod in (forms, gemini, data_engine):
        assert mod.COMPANY_PDFS is pdf_cache.COMPANY_PDFS
        assert "Q4" in mod.COMPANY_PDFS["NBHI"], f"{mod.__name__} still sees the old period"


def test_pdfs_arriving_after_set_period_are_picked_up(tmp_path, monkeypatch):
    """Phase 1 downloads into the period's directory AFTER the period is set,
    so the refresh must re-scan the filesystem rather than cache on the
    period alone."""
    from competitor_analysis import paths
    from competitor_analysis.extraction import pdf_cache

    _seed(tmp_path, "Q4", ["NBHI"])
    monkeypatch.setattr(paths, "DOWNLOADS_DIR", tmp_path)

    cfg.set_period("FY25-26", "Q4")
    assert set(pdf_cache.COMPANY_PDFS) == {"NBHI"}

    _seed(tmp_path, "Q4", ["Star Health"])          # lands mid-run
    pdf_cache.refresh_company_pdfs()
    assert set(pdf_cache.COMPANY_PDFS) == {"NBHI", "Star Health"}


def test_reverse_lookup_drops_the_old_period(tmp_path, monkeypatch):
    """_PDF_TO_COMPANY must be rebuilt alongside COMPANY_PDFS, not merely
    added to: it names the company an uncached parse is filed under, so a
    surviving entry for the previous period's path keeps that file
    attributable after the switch."""
    from competitor_analysis import paths
    from competitor_analysis.extraction import pdf_cache
    import os

    _seed(tmp_path, "Q3", ["NBHI"])
    _seed(tmp_path, "Q4", ["NBHI"])
    monkeypatch.setattr(paths, "DOWNLOADS_DIR", tmp_path)

    cfg.set_period("FY25-26", "Q3")
    q3_path = pdf_cache.COMPANY_PDFS["NBHI"]
    cfg.set_period("FY25-26", "Q4")
    q4_path = pdf_cache.COMPANY_PDFS["NBHI"]

    assert q3_path != q4_path
    assert pdf_cache._company_for_path(q4_path) == "NBHI"
    assert os.path.normpath(q3_path) not in pdf_cache._PDF_TO_COMPANY


def test_value_column_labels_follow_the_period():
    """sync_period_headers exists to stop a run presenting its figures under
    the wrong quarter's heading. Its labels were resolved once at import, so
    in the long-lived API server it wrote the FIRST run's period onto every
    later run's workbook - defeated by its own stale input."""
    import openpyxl
    from competitor_analysis.extraction import data_engine as p2

    cfg.set_period("FY26-27", "Q1")
    ws = openpyxl.Workbook().active
    p2.sync_period_headers(ws)

    assert ws.cell(row=1, column=p2.COL[p2.CUR]).value == "FY26-27_Q1"
    assert ws.cell(row=1, column=p2.COL[p2.PRIOR]).value == "FY25-26_Q1"
    # The positional invariant the workbook layout depends on must survive.
    assert p2.HEADERS[7] == p2.CUR_PERIOD and p2.HEADERS[8] == p2.PRIOR_PERIOD
    assert p2.COL[p2.CUR] == 8 and p2.COL[p2.PRIOR] == 9


def test_gic_workbook_path_follows_the_period(tmp_path, monkeypatch):
    """GIC_PATH was a DEFAULT ARGUMENT on GicData.__init__, bound at def
    time, and pipeline.py constructs GicData() with no argument - so every
    run after the first read the first run's quarter of GIC.xlsx."""
    import openpyxl
    from competitor_analysis import paths
    from competitor_analysis.extraction import data_engine as p2

    seen = {}

    def _capture(path, **kw):
        seen["path"] = path
        raise RuntimeError("path captured")

    monkeypatch.setattr(paths, "DOWNLOADS_DIR", tmp_path)
    monkeypatch.setattr(openpyxl, "load_workbook", _capture)

    cfg.set_period("FY25-26", "Q4")
    with pytest.raises(RuntimeError, match="path captured"):
        p2.GicData()
    assert seen["path"] == cfg.gic_path()
    assert "Q4" in seen["path"]


def test_trends_footnote_names_the_configured_period():
    from competitor_analysis.reporting import report

    cfg.set_period("FY26-27", "Q1")
    note = report._trends_note()
    assert "FY25-26 Q1" in note and "FY26-27 Q1" in note
    assert "Q3" not in note
